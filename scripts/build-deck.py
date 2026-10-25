"""Build docs/ClaimTrail_pitch.pptx from docs/PITCH_DECK.md (optional; needs `pip install python-pptx`)."""
import re, sys, pathlib
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR

root = pathlib.Path(__file__).resolve().parent.parent
src = (root / "docs" / "PITCH_DECK.md").read_text(encoding="utf-8")
parts = re.split(r"^## Slide \d+ — ", src, flags=re.M)[1:]
assert len(parts) == 7, f"expected 7 slides, found {len(parts)}"

INK, TEAL, MUTED, BG = RGBColor(0x1F, 0x29, 0x37), RGBColor(0x0B, 0x6E, 0x6A), RGBColor(0x5B, 0x65, 0x73), RGBColor(0xF6, 0xF7, 0xF9)
prs = Presentation(); prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
blank = prs.slide_layouts[6]

def clean(t): return re.sub(r"[`*]", "", t).strip()

def box(slide, x, y, w, h, text, size, bold=False, color=INK):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = tb.text_frame; tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP; p = tf.paragraphs[0]; r = p.add_run(); r.text = text
    r.font.size, r.font.bold, r.font.color.rgb, r.font.name = Pt(size), bold, color, "Calibri"; return tf

for i, part in enumerate(parts, 1):
    title, _, body = part.partition("\n")
    lines = [l for l in body.strip().splitlines()]
    headline = next((clean(l) for l in lines if l.startswith("**")), "")
    bullets = [clean(l[2:]) for l in lines if l.startswith("- ")]
    note = next((clean(l) for l in lines if l.startswith("*Speaker note")), "")
    table = [[clean(c) for c in l.strip("|").split("|")] for l in lines if l.startswith("|") and not re.match(r"^\|[-| ]+\|$", l)]
    s = prs.slides.add_slide(blank)
    s.background.fill.solid(); s.background.fill.fore_color.rgb = BG
    bar = s.shapes.add_shape(1, 0, 0, Inches(0.25), prs.slide_height); bar.fill.solid(); bar.fill.fore_color.rgb = TEAL; bar.line.fill.background()
    box(s, 0.7, 0.4, 11.5, 0.6, f"{i:02d} · {clean(title)}", 16, True, TEAL)
    box(s, 0.7, 1.0, 11.9, 1.4, headline, 32, True)
    y = 2.6
    if table:
        rows, cols = len(table), len(table[0])
        t = s.shapes.add_table(rows, cols, Inches(0.7), Inches(y), Inches(11.9), Inches(2.6)).table
        for r_i, row in enumerate(table):
            for c_i, cell in enumerate(row):
                c = t.cell(r_i, c_i); c.text = cell; c.fill.solid()
                c.fill.fore_color.rgb = TEAL if r_i == 0 else RGBColor(0xFF, 0xFF, 0xFF)
                for p in t.cell(r_i, c_i).text_frame.paragraphs:
                    for r in p.runs: r.font.size = Pt(15); r.font.bold = r_i == 0; r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF) if r_i == 0 else INK
        y += 3.0
    if bullets:
        tf = box(s, 0.7, y, 11.9, 7.0 - y, "• " + bullets[0], 20)
        for b in bullets[1:]:
            p = tf.add_paragraph(); p.space_before = Pt(8); r = p.add_run(); r.text = "• " + b; r.font.size = Pt(20); r.font.color.rgb = INK; r.font.name = "Calibri"
    if note: s.notes_slide.notes_text_frame.text = note.replace("Speaker note:", "").strip()
    box(s, 0.7, 7.0, 11.9, 0.3, "ClaimTrail · hackathon MVP · synthetic demo data only · not a claim decision tool", 10, False, MUTED)

out = root / "docs" / "ClaimTrail_pitch.pptx"; prs.save(out); print("wrote", out, "slides:", len(prs.slides))
